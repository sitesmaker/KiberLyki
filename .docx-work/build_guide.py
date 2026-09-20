from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"D:\Projects\KiberLyki\docs\KiberLyki_Инструкция_по_функционалу.docx"
BLUE = "163A5F"
PALE_BLUE = "EEF4F8"
LIGHT_GRAY = "D9D9D9"
TEXT = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), LIGHT_GRAY)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Arial", size=10.5, bold=False, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def style_paragraph_runs(paragraph, size=10.5, bold=False, color=TEXT):
    for run in paragraph.runs:
        set_font(run, size=size, bold=bold, color=color)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_font(first, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_font(run)
    return p


STEP_NUMBER = 0


def reset_numbers():
    global STEP_NUMBER
    STEP_NUMBER = 0


def add_number(doc, text, level=0):
    global STEP_NUMBER
    STEP_NUMBER += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28 + level * 0.22)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    number = p.add_run(f"{STEP_NUMBER}.  ")
    set_font(number)
    run = p.add_run(text)
    set_font(run)
    return p


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        element = p_bdr.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            p_bdr.append(element)
        element.set(qn("w:val"), "nil")


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F3F3")
    p_pr.append(shd)
    for index, line in enumerate(lines):
        run = p.add_run(line)
        set_font(run, name="Consolas", size=9.3)
        if index != len(lines) - 1:
            run.add_break()
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    for idx, heading in enumerate(headers):
        cell = header.cells[idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(heading)
        set_font(run, size=9.5, bold=True, color=RGBColor(255, 255, 255))
        if widths:
            cell.width = Inches(widths[idx])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_font(run, size=9.2, bold=(idx == 0))
            if widths:
                cell.width = Inches(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(27)
title_style.font.bold = True
title_style.font.color.rgb = TEXT
title_style.paragraph_format.space_after = Pt(12)

for style_name, size, before, after in (("Heading 1", 18, 14, 7), ("Heading 2", 13.5, 11, 5), ("Heading 3", 11.5, 8, 4)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = TEXT
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("Инструкция по функционалу KiberLyki")
remove_paragraph_borders(p)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
run = p.add_run("React и Strapi 5")
set_font(run, size=15, bold=True, color=RGBColor(50, 50, 50))

add_body(doc, "Документ объясняет, какие данные заполнить в Strapi и как последовательно проверить личный кабинет, профили игроков, управление командой, заявки на турниры, работу организатора и турнирную сетку.")
add_body(doc, "Основной принцип: Strapi хранит данные и выполняет серверные проверки, а React показывает страницы и отправляет действия пользователя в API.", bold_lead="Основной принцип:")

doc.add_paragraph().paragraph_format.space_after = Pt(10)
add_table(doc, ["Часть системы", "Адрес", "Назначение"], [
    ("Сайт", "http://localhost:5173", "Страницы для гостей, игроков, капитанов и организаторов"),
    ("Strapi", "http://localhost:1337/admin", "Создание контента, турниров, дисциплин и назначение ролей"),
], widths=[1.25, 2.05, 3.55])

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
run = p.add_run("Актуально для текущей учебной версии проекта")
set_font(run, size=10, bold=True, color=RGBColor(70, 70, 70))

doc.add_page_break()

# Contents
doc.add_heading("Содержание", level=1)
for item in [
    "1 Запуск проекта",
    "2 Роли и права пользователей",
    "3 Подготовка данных в Strapi",
    "4 Проверка профиля игрока",
    "5 Проверка личного кабинета капитана",
    "6 Регистрация команды на турнир",
    "7 Работа организатора и турнирная сетка",
    "8 Справочник статусов",
    "9 Полный сценарий проверки",
    "10 Ограничения текущей версии",
]:
    add_bullet(doc, item)

doc.add_heading("1 Запуск проекта", level=1)
add_body(doc, "Backend и frontend запускаются в двух отдельных терминалах. При первом запуске установите зависимости в обеих папках.")

doc.add_heading("Backend", level=2)
add_code(doc, [
    r"cd D:\Projects\KiberLyki\backend",
    "npm install",
    "npm run dev",
])
add_body(doc, "После запуска откройте панель управления Strapi по адресу http://localhost:1337/admin.")

doc.add_heading("Frontend", level=2)
add_code(doc, [
    r"cd D:\Projects\KiberLyki\client",
    "npm install",
    "Copy-Item .env.example .env",
    "npm run dev",
])
add_body(doc, "Vite напечатает адрес сайта в терминале. Обычно это http://localhost:5173. Переменная VITE_API_URL в файле .env должна указывать на http://localhost:1337.")

doc.add_heading("После изменения схем Strapi", level=2)
add_body(doc, "Полностью остановите backend сочетанием Ctrl+C и снова выполните npm run dev. Простого обновления страницы недостаточно, если менялись поля content type.")

doc.add_heading("2 Роли и права пользователей", level=1)
add_table(doc, ["Роль", "Где используется", "Что может делать"], [
    ("Public", "Сайт без входа", "Смотреть опубликованные новости, команды и турниры"),
    ("Authenticated", "React сайт", "Заполнять профиль, состоять в команде, а при создании команды стать капитаном"),
    ("Organizer", "React сайт", "Рассматривать заявки, формировать сетку и вводить результаты матчей"),
    ("Strapi Admin", "Панель Strapi", "Создавать контент и управлять данными через административную панель"),
], widths=[1.35, 1.55, 3.95])
add_body(doc, "Strapi Admin и Organizer — разные аккаунты. Администратор панели не может автоматически войти на сайт как организатор. Для сайта нужен отдельный пользователь Users and Permissions с ролью Organizer.", bold_lead="Важно:")

doc.add_heading("3 Подготовка данных в Strapi", level=1)

doc.add_heading("Дисциплины", level=2)
add_body(doc, "Откройте Content Manager, затем Discipline и создайте игровые дисциплины.")
add_table(doc, ["Поле", "Пример", "Для чего нужно"], [
    ("name", "Dota 2", "Название дисциплины на сайте"),
    ("slug", "dota-2", "Уникальный технический адрес"),
    ("teamSize", "5", "Рекомендуемый размер состава для дисциплины"),
    ("icon", "Изображение", "Необязательная иконка"),
], widths=[1.3, 1.65, 3.9])
add_body(doc, "Создайте как минимум две дисциплины, например Dota 2 и Counter Strike. Тогда в кабинете можно проверить мультивыбор дисциплин для одной команды.")

doc.add_heading("Турнир", level=2)
add_body(doc, "Откройте Content Manager, затем Tournament и заполните поля ниже.")
add_table(doc, ["Поле", "Что указать", "Пояснение"], [
    ("title", "Турнир по Counter Strike", "Название карточки"),
    ("slug", "cs-autumn-cup", "Уникальное значение без пробелов"),
    ("description", "Описание турнира", "Текст на детальной странице"),
    ("rules", "Правила участия", "Допустимые правила и ограничения"),
    ("cover", "Обложка", "Необязательное изображение"),
    ("format", "single_elimination", "Олимпийская система на выбывание"),
    ("teamSize", "5", "Сколько игроков капитан обязан выбрать"),
    ("maxTeams", "8", "Максимальное число активных заявок"),
    ("registrationStartsAt", "Дата в прошлом", "Начало окна регистрации"),
    ("registrationEndsAt", "Дата в будущем", "Конец окна регистрации"),
    ("startsAt", "После конца регистрации", "Планируемое начало турнира"),
    ("phase", "registration", "Текущая бизнес фаза турнира"),
    ("discipline", "Counter Strike", "Одна дисциплина конкретного турнира"),
], widths=[1.65, 1.85, 3.35])
add_body(doc, "После сохранения нажмите Publish. Для возможности подать заявку должны одновременно выполняться три условия: phase равен registration, текущая дата находится внутри окна регистрации и запись опубликована.", bold_lead="Для открытия регистрации:")

doc.add_heading("Организатор", level=2)
reset_numbers()
add_number(doc, "Зарегистрируйте обычный аккаунт на странице /login или создайте пользователя через Content Manager, раздел User.")
add_number(doc, "В Strapi откройте этого пользователя.")
add_number(doc, "Установите Role равной Organizer, Confirmed равной true, Blocked равной false.")
add_number(doc, "Сохраните пользователя и войдите этим аккаунтом на React сайте.")

doc.add_heading("Новости и главная страница", level=2)
add_table(doc, ["Сущность", "Основные поля", "Где видна"], [
    ("Article", "title, slug, description, content, cover, gallery, author", "Список новостей и детальная страница"),
    ("First screen", "title, content, media", "Первый экран главной страницы"),
    ("About", "title, description, image", "Блок о проекте"),
    ("Staff", "имя, фамилия, специализация, описание, фото, дата рождения", "Блок сотрудников"),
], widths=[1.35, 3.25, 2.25])
add_body(doc, "Для сущностей с Draft and Publish после Save обязательно нажимайте Publish. Иначе публичный API не вернёт запись.")

doc.add_page_break()
doc.add_heading("4 Проверка профиля игрока", level=1)
reset_numbers()
add_number(doc, "Войдите на сайт через /login.")
add_number(doc, "Откройте Личный кабинет и нажмите Мой профиль.")
add_number(doc, "Заполните никнейм, страну, имя, фамилию и текст О себе.")
add_number(doc, "При необходимости загрузите аватар и нажмите Сохранить профиль.")
add_number(doc, "Если аккаунт создал капитан, в форме Сменить пароль укажите временный и новый пароль.")
add_body(doc, "Профиль принадлежит самому игроку. Капитан управляет участием в команде, но не редактирует личные данные игрока.", bold_lead="Разделение ответственности:")

doc.add_heading("5 Проверка личного кабинета капитана", level=1)

doc.add_heading("Создание команды", level=2)
reset_numbers()
add_number(doc, "Войдите пользователем, который пока не состоит в активной команде.")
add_number(doc, "Откройте /cabinet.")
add_number(doc, "Укажите название и описание команды.")
add_number(doc, "Отметьте несколько дисциплин, например Dota 2 и Counter Strike.")
add_number(doc, "Нажмите Создать команду. Создавший пользователь становится капитаном и основным игроком.")

doc.add_heading("Добавление нового игрока", level=2)
add_body(doc, "Форма Создать нового игрока нужна, когда у человека ещё нет аккаунта. Заполните username, никнейм, email, временный пароль и позицию. Backend создаст аккаунт, профиль и активное участие в команде.")

doc.add_heading("Добавление зарегистрированного игрока", level=2)
reset_numbers()
add_number(doc, "Попросите игрока самостоятельно зарегистрироваться на /login.")
add_number(doc, "Войдите аккаунтом капитана и откройте кабинет.")
add_number(doc, "В блоке Добавить зарегистрированного введите точный username или email.")
add_number(doc, "Выберите позицию и нажмите Добавить в состав.")
add_body(doc, "Пользователя нельзя добавить, если он уже состоит в другой активной команде.")

doc.add_heading("Позиции в команде", level=2)
add_table(doc, ["Значение", "Отображение", "Участие в заявке"], [
    ("main", "Основной", "Можно выбрать в состав на турнир"),
    ("substitute", "Запасной", "Можно выбрать в состав на турнир"),
    ("coach", "Тренер", "Не считается игроком для турнирной заявки"),
], widths=[1.35, 2.0, 3.5])

doc.add_heading("Удаление и передача капитанства", level=2)
add_body(doc, "Капитан может удалить из состава любого активного участника, кроме самого себя. Это мягкое удаление: аккаунт, профиль и история заявок сохраняются.")
add_body(doc, "Для передачи прав выберите активного игрока в блоке Передать капитанство. После операции новый капитан получает элементы управления, а прежний остаётся обычным участником.")
add_body(doc, "Если капитан хочет уйти, сначала передайте права другому игроку, затем новый капитан сможет удалить прежнего.", bold_lead="Правильный порядок ухода капитана:")

doc.add_heading("6 Регистрация команды на турнир", level=1)
reset_numbers()
add_number(doc, "Убедитесь, что команда поддерживает дисциплину турнира.")
add_number(doc, "Проверьте, что phase турнира равна registration и окно дат открыто.")
add_number(doc, "Откройте /tournaments и перейдите в детальную карточку.")
add_number(doc, "Выберите ровно столько активных игроков, сколько указано в tournament.teamSize. Тренеры недоступны для выбора.")
add_number(doc, "Отправьте заявку. Она получит registrationStatus pending.")
add_body(doc, "Одна команда может выбрать несколько дисциплин и подать отдельную заявку в турнир по каждой из них. Состав выбирается заново для каждого турнира, поэтому игроки в заявке по Dota 2 и Counter Strike могут отличаться.")
add_body(doc, "Повторную активную заявку той же команды в тот же турнир создать нельзя. Капитан может отозвать pending или approved заявку до начала турнира.")

doc.add_heading("7 Работа организатора и турнирная сетка", level=1)
reset_numbers()
add_number(doc, "Войдите аккаунтом с ролью Organizer.")
add_number(doc, "Откройте детальную страницу турнира. Появится блок Управление турниром.")
add_number(doc, "Рассмотрите заявки и нажмите Принять или Отклонить.")
add_number(doc, "Одобрите минимум две команды.")
add_number(doc, "Нажмите Сформировать сетку. Повторно формировать существующую сетку нельзя.")
add_number(doc, "Введите разные неотрицательные результаты команд и сохраните матч.")
add_body(doc, "Победитель автоматически переходит в следующий матч. Если число команд не является степенью двойки, система создаёт свободные места bye и автоматически продвигает соответствующие команды.")
add_body(doc, "После создания сетки phase турнира автоматически становится active. После сохранения результата финала phase становится finished.")

doc.add_page_break()
doc.add_heading("8 Справочник статусов", level=1)
add_body(doc, "В проекте используются разные наборы состояний. Их нельзя смешивать: phase описывает турнир, registrationStatus описывает заявку, membershipStatus описывает участие игрока, а matchStatus описывает матч.")

doc.add_heading("Публикация Strapi", level=2)
add_table(doc, ["Значение", "Что означает", "Как меняется"], [
    ("draft", "Запись сохранена как черновик и обычно не видна публичному API", "Кнопка Save или Unpublish в Strapi"),
    ("published", "Запись опубликована и доступна сайту", "Кнопка Publish в Strapi"),
], widths=[1.35, 3.25, 2.25])
add_body(doc, "Это системный статус публикации Strapi. Он не заменяет поле phase турнира.")

doc.add_heading("Фаза турнира phase", level=2)
add_table(doc, ["Значение", "Назначение", "Текущее поведение"], [
    ("draft", "Турнир готовится", "Заявка с сайта запрещена"),
    ("registration", "Открыт приём заявок", "Заявка разрешена только внутри заданного окна дат"),
    ("check_in", "Подтверждение участников перед стартом", "Сейчас используется как информационная фаза; отдельный check in ещё не реализован"),
    ("active", "Турнир начался", "Устанавливается автоматически после создания сетки"),
    ("finished", "Турнир завершён", "Устанавливается автоматически после результата финала"),
    ("cancelled", "Турнир отменён", "Устанавливается вручную в текущей версии"),
], widths=[1.35, 2.35, 3.15])

doc.add_heading("Статус заявки registrationStatus", level=2)
add_table(doc, ["Значение", "Что означает", "Кто устанавливает"], [
    ("pending", "Заявка отправлена и ожидает решения", "Система при регистрации команды"),
    ("approved", "Команда допущена к турниру", "Организатор кнопкой Принять"),
    ("rejected", "Организатор отказал команде", "Организатор кнопкой Отклонить"),
    ("withdrawn", "Капитан отозвал заявку", "Капитан кнопкой Отозвать"),
], widths=[1.45, 3.25, 2.15])
add_body(doc, "Withdrawn не удаляет запись. Заявка остаётся в истории, но не участвует в лимите и формировании сетки.")

doc.add_heading("Статус участия membershipStatus", level=2)
add_table(doc, ["Значение", "Что означает", "Текущее использование"], [
    ("invited", "Игрок приглашён, но ещё не подтвердил участие", "Зарезервировано; приглашения пока не реализованы"),
    ("active", "Игрок состоит в команде", "Используется в составе и доступен для выбора, кроме тренера"),
    ("left", "Игрок самостоятельно покинул команду", "Зарезервировано; самостоятельный выход пока не реализован"),
    ("removed", "Капитан удалил игрока", "Устанавливается при удалении из состава"),
], widths=[1.35, 2.8, 2.7])

doc.add_heading("Статус матча matchStatus", level=2)
add_table(doc, ["Значение", "Что означает", "Текущее использование"], [
    ("pending", "Матч создан и ждёт результата", "Начальное состояние матча сетки"),
    ("scheduled", "Матчу назначено время", "Зарезервировано; отдельное планирование пока не реализовано"),
    ("live", "Матч идёт", "Зарезервировано; live режим пока не реализован"),
    ("finished", "Результат матча зафиксирован", "Система после ввода счёта или автоматического bye"),
], widths=[1.35, 2.8, 2.7])

doc.add_page_break()
doc.add_heading("9 Полный сценарий проверки", level=1)
add_body(doc, "Для проверки всего турнирного цикла нужны организатор и минимум две команды. Для реалистичного турнира 5 на 5 потребуется десять игровых аккаунтов. Для быстрой учебной проверки можно создать отдельный тестовый турнир с teamSize равным 2.")
reset_numbers()
for text in [
    "Создайте две дисциплины в Strapi.",
    "Создайте и опубликуйте турнир с phase registration и актуальными датами.",
    "Создайте аккаунт организатора и назначьте роль Organizer.",
    "Зарегистрируйте первого капитана, создайте команду и выберите обе дисциплины.",
    "Создайте игрока с временным паролем и добавьте ещё одного заранее зарегистрированного игрока.",
    "Войдите игроком, заполните профиль и смените временный пароль.",
    "Создайте второго капитана и вторую команду с достаточным составом.",
    "Подайте заявки обеих команд на один турнир.",
    "Войдите организатором, одобрите обе заявки и сформируйте сетку.",
    "Введите результат финала и проверьте phase finished.",
    "Вернитесь капитаном, проверьте удаление игрока и передачу капитанства.",
    "Создайте второй турнир по другой дисциплине и убедитесь, что та же команда может подать новую заявку.",
]:
    add_number(doc, text)

doc.add_heading("10 Ограничения текущей версии", level=1)
add_bullet(doc, "Приглашения со статусом invited и самостоятельный выход left описаны в модели, но отдельные пользовательские сценарии ещё не реализованы.")
add_bullet(doc, "Фаза check_in и состояния scheduled и live пока информационные и не имеют отдельных экранов управления.")
add_bullet(doc, "Название, описание и логотип команды можно хранить в Strapi, но текущий кабинет капитана редактирует только дисциплины после создания команды.")
add_bullet(doc, "Сетка поддерживает формат single elimination. Другие форматы турниров пока не реализованы.")
add_bullet(doc, "Нельзя вводить ничью: результат матча должен определить победителя.")

doc.add_heading("Контрольный список", level=2)
for text in [
    "Backend и frontend запущены без ошибок",
    "Турнир опубликован и имеет phase registration",
    "Окно регистрации включает текущую дату",
    "Команда поддерживает дисциплину турнира",
    "В заявке выбрано ровно tournament.teamSize игроков",
    "Есть минимум две approved заявки",
    "Организатор видит панель управления турниром",
    "После финала турнир имеет phase finished",
]:
    add_bullet(doc, "☐ " + text)

# Apply keep-with-next to headings and consistent body color.
for paragraph in doc.paragraphs:
    if paragraph.style and paragraph.style.name.startswith("Heading"):
        set_keep_with_next(paragraph)
    for run in paragraph.runs:
        if run.font.color.rgb is None:
            run.font.color.rgb = TEXT

doc.core_properties.title = "Инструкция по функционалу KiberLyki"
doc.core_properties.subject = "Запуск, настройка Strapi и проверка функций React приложения"
doc.core_properties.author = "KiberLyki"
doc.core_properties.keywords = "React, Strapi, турниры, команды, инструкция"
doc.save(OUTPUT)
print(OUTPUT)
